# -*- coding: utf-8 -*-
"""
将 design_proposal.md 转换为格式规范、排版优美的 Word 文档 (.docx)
确保符合信创大赛评审匿名与专业规范
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

MD_PATH = r"D:\Dream\bearpi-hm_nano\design_proposal.md"
DOCX_PATH = r"D:\Dream\bearpi-hm_nano\独居老人居家安全监测与应急辅助系统_设计方案.docx"

def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # 页面边距
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # 默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '微软雅黑'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    with open(MD_PATH, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        # 过滤分割线
        clean_rows = []
        for r in table_rows:
            if re.match(r'^[\|\s\-:]+$', r.strip()):
                continue
            cells = [c.strip() for c in r.strip().strip('|').split('|')]
            clean_rows.append(cells)

        if not clean_rows:
            in_table = False
            table_rows = []
            return

        cols_count = max(len(r) for r in clean_rows)
        tbl = doc.add_table(rows=len(clean_rows), cols=cols_count)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        for row_idx, row_data in enumerate(clean_rows):
            is_header = (row_idx == 0)
            for col_idx in range(cols_count):
                val = row_data[col_idx] if col_idx < len(row_data) else ""
                cell = tbl.cell(row_idx, col_idx)
                cell.text = val
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                
                # 单元格字体
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    if is_header:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in p.runs:
                            run.font.bold = True
                            run.font.size = Pt(10.5)
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    else:
                        for run in p.runs:
                            run.font.size = Pt(10)
                            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

                if is_header:
                    set_cell_background(cell, "1F497D")
                elif row_idx % 2 == 1:
                    set_cell_background(cell, "F2F5F9")
                else:
                    set_cell_background(cell, "FFFFFF")

        # 添加空行
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(6)

        in_table = False
        table_rows = []

    def flush_code():
        nonlocal in_code_block, code_lines
        if code_lines:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.left_indent = Inches(0.2)
            run = p.add_run('\n'.join(code_lines))
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        in_code_block = False
        code_lines = []

    for line in lines:
        stripped = line.rstrip()

        # 代码块判定
        if stripped.startswith('```'):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(stripped)
            continue

        # 表格判定
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_rows.append(stripped)
            continue
        elif in_table:
            flush_table()

        # 分割线
        if stripped == '---':
            continue

        # 标题处理
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(12)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped[2:].strip())
            run.font.bold = True
            run.font.size = Pt(20)
            run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
            continue
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(stripped[3:].strip())
            run.font.bold = True
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
            continue
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[4:].strip())
            run.font.bold = True
            run.font.size = Pt(12.5)
            run.font.color.rgb = RGBColor(0x36, 0x60, 0x92)
            continue

        # 列表处理
        if stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.2
            text = stripped[2:].strip()
            # 简单粗体解析 **text**
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
            continue

        # 数字列表
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.2
            text = num_match.group(2).strip()
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
            continue

        # 普通段落
        if stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.first_line_indent = Inches(0.28) # 首行缩进两个汉字
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)

    if in_table:
        flush_table()
    if in_code_block:
        flush_code()

    doc.save(DOCX_PATH)
    print(f"[SUCCESS] Word 文档已生成: {DOCX_PATH}")

if __name__ == '__main__':
    create_document()
