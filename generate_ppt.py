# -*- coding: utf-8 -*-
import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation(image_path="D:/Dream/elderly_care_keynote.jpg", output_path="D:/Dream/独居老人居家安全监测系统_路演答辩PPT.pptx"):
    prs = Presentation()
    # 16:9 Widescreen standard
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette - High-tech Dark Theme
    C_BG = RGBColor(11, 17, 30)          # #0B111E 深蓝黑
    C_HEADER = RGBColor(16, 25, 44)      # #10192C 导航色
    C_CARD = RGBColor(21, 30, 50)        # #151E32 卡片底色
    C_CARD_ALT = RGBColor(25, 36, 60)    # #19243C 稍亮卡片
    C_BORDER = RGBColor(35, 49, 78)      # #23314E 边框
    C_PRIMARY = RGBColor(0, 125, 255)    # #007DFF 鸿蒙蓝
    C_CYAN = RGBColor(0, 210, 255)       # #00D2FF 科技青
    C_SUCCESS = RGBColor(0, 200, 83)     # #00C853 绿色
    C_WARNING = RGBColor(255, 171, 0)    # #FFAB00 警告橙
    C_DANGER = RGBColor(255, 59, 48)     # #FF3B30 报警红
    C_TEXT_MAIN = RGBColor(240, 244, 248)# 主文字白
    C_TEXT_SUB = RGBColor(141, 164, 190) # 辅助灰蓝
    C_MUTED = RGBColor(90, 110, 135)     # 暗调文字

    def set_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title, category="赛题四 · 独居老人居家安全监测与应急辅助系统", step_tag="OpenHarmony 架构"):
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.1))
        header.fill.solid()
        header.fill.fore_color.rgb = C_HEADER
        header.line.color.rgb = C_BORDER
        header.line.width = Pt(1)

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(9.8), Inches(0.85))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = category
        p0.font.name = "Microsoft YaHei"
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = C_PRIMARY

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = "Microsoft YaHei"
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = C_TEXT_MAIN

        # Right badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.6), Inches(0.32), Inches(2.0), Inches(0.44))
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor(0, 40, 85)
        badge.line.color.rgb = C_PRIMARY
        badge.line.width = Pt(1)
        btf = badge.text_frame
        btf.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = 0
        bp = btf.paragraphs[0]
        bp.text = step_tag
        bp.alignment = PP_ALIGN.CENTER
        bp.font.name = "Microsoft YaHei"
        bp.font.size = Pt(11)
        bp.font.bold = True
        bp.font.color.rgb = C_CYAN

    def add_card(slide, left, top, width, height, fill_color=C_CARD, border_color=C_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # =========================================================================
    # SLIDE 1: 封面页 (Title Slide)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1)
    add_card(s1, Inches(0.8), Inches(0.9), Inches(11.733), Inches(5.7), fill_color=C_HEADER, border_color=C_BORDER)

    tb1 = s1.shapes.add_textbox(Inches(1.3), Inches(1.4), Inches(6.8), Inches(4.7))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "第三届河南省高校信息技术应用创新大赛 · 开源鸿蒙赛项"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_CYAN
    p.space_after = Pt(14)

    p = tf1.add_paragraph()
    p.text = "独居老人居家安全监测\n与应急辅助系统"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_MAIN
    p.space_after = Pt(14)

    p = tf1.add_paragraph()
    p.text = "基于 OpenHarmony 的软硬件南北向一体化健康守护平台"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(15)
    p.font.color.rgb = C_PRIMARY
    p.space_after = Pt(22)

    p = tf1.add_paragraph()
    p.text = "🛡️ 无感被动感知  |  🚨 毫秒级声光抢占  |  📋 社区闭环台账  |  📊 规律健康图谱"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_SUB
    p.space_after = Pt(26)

    p = tf1.add_paragraph()
    p.text = "【参赛赛题】 赛题四 · 南北结合软硬件开发类  （作品匿名汇报）"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(11)
    p.font.color.rgb = C_MUTED

    if os.path.exists(image_path):
        s1.shapes.add_picture(image_path, Inches(8.3), Inches(1.3), Inches(3.9), Inches(4.9))

    # =========================================================================
    # SLIDE 2: 行业痛点与攻关目标 (Pain Points & Vision)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "行业痛点与攻关目标：为什么要做“无感守护”？")

    # 3 Pain Point Cards
    cards_data = [
        ("传统智能手环 / 穿戴", "频繁充电、容易遗忘", "老年人常忘充忘戴，沐浴、睡眠等高危跌倒场景往往主动摘卸，导致监护“窗口期”频繁失效。", C_DANGER),
        ("室内摄像头安防监控", "严重侵犯隐私与尊严", "在卫生间、卧室等跌倒高发场所，老人极度抗拒摄像头；且面临严重的家庭隐私泄露安全隐患。", C_WARNING),
        ("常规物联网预警设备", "“只报不管”，缺乏处置", "报警信息发到云端就草草结束，缺乏网格员与子女的排查确认与处置记录，无法形成闭环安全台账。", C_PRIMARY)
    ]
    for idx, (title, sub, desc, color) in enumerate(cards_data):
        left = Inches(0.8 + idx * 3.98)
        card = add_card(s2, left, Inches(1.4), Inches(3.78), Inches(3.8))
        
        # Color bar
        bar = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.4), Inches(3.78), Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        tb = s2.shapes.add_textbox(left + Inches(0.25), Inches(1.7), Inches(3.28), Inches(3.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"痛点 0{idx+1} · {title}"
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_MAIN
        p.space_after = Pt(8)

        p = tf.add_paragraph()
        p.text = sub
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(13)
        p.font.color.rgb = color
        p.space_after = Pt(14)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(12)
        p.font.color.rgb = C_TEXT_SUB
        p.space_after = Pt(10)

    # Bottom summary card
    bot_card = add_card(s2, Inches(0.8), Inches(5.4), Inches(11.733), Inches(1.5), fill_color=C_HEADER)
    btb = s2.shapes.add_textbox(Inches(1.1), Inches(5.5), Inches(11.1), Inches(1.3))
    btf = btb.text_frame
    btf.word_wrap = True
    bp = btf.paragraphs[0]
    bp.text = "💡 本项目核心破局之道 ——“非感官接触式、高隐私保护、全流程闭环”"
    bp.font.name = "Microsoft YaHei"
    bp.font.size = Pt(15)
    bp.font.bold = True
    bp.font.color.rgb = C_CYAN
    bp.space_after = Pt(6)

    bp = btf.add_paragraph()
    bp.text = "依托开源鸿蒙（OpenHarmony）轻量系统与微型传感矩阵，无需老人佩戴任何设备，不采集任何视频图像，在卧室与卫生间提供全天候无死角的隐形安全守护，并深度打通“预警-排查-处置-归档”的社区闭环链路。"
    bp.font.name = "Microsoft YaHei"
    bp.font.size = Pt(12)
    bp.font.color.rgb = C_TEXT_MAIN

    # =========================================================================
    # SLIDE 3: 系统总体技术架构 (System Architecture)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "系统总体技术架构：软硬协同 · 南北贯通", step_tag="端 - 管 - 云 - 屏")

    arch_layers = [
        ("北向综合控制与闭环管理层", "OpenHarmony 响应式架构 · 实时监控仪表盘 + 事件闭环处置工作台 + 24h生活规律图谱", [
            "大屏态势感知：高分贝声光同步、毫秒级倒计时刷新",
            "事件处置台账：子女/网格员上门核实，一键消警与电子归档",
            "规律大数据：24小时活跃时段柱状图、7日活动衰退趋势预警"
        ], C_PRIMARY),
        ("通信传输与多模调度层", "双轨并发遥测流 · 局域网极速通信 (<50ms) + 华为云物联网平台 (IoTDA)", [
            "实时高速遥测：Web Serial / UART 标准 JSON 遥测流（答辩演示首选，0延迟不死机）",
            "物联网云中转：MQTT 协议栈属性上报与远程下发控制，无缝对接华为智能终端"
        ], C_CYAN),
        ("南向设备感知与控制层", "OpenHarmony LiteOS-M 内核 · 海思 Hi3861 RISC-V 主控 + 微型传感集群", [
            "人体活动感应：E53_IS1 热释电红外（GPIO_7 中断，实时刷新最后活动时间）",
            "环境微气候：SHT30 工业级温湿度（I²C1 总线，极温与火灾早期预警）",
            "门窗磁与求助：磁敏干簧管（GPIO_13）+ 板载 SOS 紧急求助键（GPIO_11）",
            "本地声光联动：PWM1 高分贝蜂鸣报警（GPIO_8）+ 板载黄色 LED 急促爆闪"
        ], C_SUCCESS)
    ]

    for idx, (layer_title, sub, items, col) in enumerate(arch_layers):
        top = Inches(1.4 + idx * 1.85)
        card = add_card(s3, Inches(0.8), top, Inches(11.733), Inches(1.7))

        bar = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(0.16), Inches(1.7))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tb = s3.shapes.add_textbox(Inches(1.2), top + Inches(0.15), Inches(11.1), Inches(1.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"{layer_title}  |  {sub}"
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(6)

        for item in items:
            p = tf.add_paragraph()
            p.text = "• " + item
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(11.5)
            p.font.color.rgb = C_TEXT_MAIN
            p.space_after = Pt(3)

    # =========================================================================
    # SLIDE 4: 硬件选型与巧妙的抗冲突电气设计 (Hardware Design)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "硬件选型与巧妙的抗冲突电气设计", step_tag="硬核工程技术")

    # Left card: Hardware Specs
    c_left = add_card(s4, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.5))
    tb_l = s4.shapes.add_textbox(Inches(1.1), Inches(1.6), Inches(5.0), Inches(5.1))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "⚙️ 核心硬件清单与引脚账目"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_CYAN
    p.space_after = Pt(12)

    hw_items = [
        ("主控开发板", "BearPi-HM Nano（Hi3861 RISC-V 32位，160MHz）"),
        ("人体红外探头", "E53_IS1（GPIO_7，上升沿硬件中断，感知人动）"),
        ("声光报警模块", "E53_IS1 蜂鸣器（GPIO_8 / PWM1，4kHz高分贝警报）"),
        ("微气候温湿度", "SHT30（I²C1：SDA=GPIO_0, SCL=GPIO_1）"),
        ("门窗磁传感器", "高灵敏磁敏干簧管（GPIO_13，内部强上拉）"),
        ("紧急求助按键", "板载微动 F1 按键（GPIO_11，一键抢占式呼救）"),
        ("消警复位按键", "板载微动 F2 按键（GPIO_12，现场排查恢复）")
    ]
    for name, detail in hw_items:
        p = tf_l.add_paragraph()
        p.text = f"• {name}：{detail}"
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(12)
        p.font.color.rgb = C_TEXT_MAIN
        p.space_after = Pt(6)

    # Right card: The GPIO_8 conflict resolution (Huge highlight!)
    c_right = add_card(s4, Inches(6.8), Inches(1.4), Inches(5.733), Inches(5.5), fill_color=C_HEADER)
    tb_r = s4.shapes.add_textbox(Inches(7.1), Inches(1.6), Inches(5.1), Inches(5.1))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "🔥 攻克技术难题：单E53插槽与GPIO8双占冲突"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_WARNING
    p.space_after = Pt(12)

    p = tf_r.add_paragraph()
    p.text = "【电气冲突根源】"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_DANGER
    p.space_after = Pt(4)

    p = tf_r.add_paragraph()
    p.text = "小熊派只有 1 个 E53 插座（J1）。E53_IS1 占用 GPIO_8 发送 PWM 驱动蜂鸣器；E53_IA1 同样占用 GPIO_8 驱动电机。若双板盲目共用，后初始化的驱动必将冲掉复用模式！"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(11.5)
    p.font.color.rgb = C_TEXT_SUB
    p.space_after = Pt(12)

    p = tf_r.add_paragraph()
    p.text = "【创新破解解法（免焊接 · 零冲突）】"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_SUCCESS
    p.space_after = Pt(4)

    solutions = [
        "1. 业务逻辑剪裁：老人居家监护需要蜂鸣器，压根不需要浇花电机！彻底废除电机驱动。",
        "2. 结构主从配置：E53_IS1 主板原位直插 J1，保障红外与蜂鸣器高频信号零接触电阻。",
        "3. 总线飞线引出：E53_IA1 仅借用 I²C1 总线（4根杜邦线连3.3V/GND/SCL/SDA），空出GPIO8引脚。",
        "4. 最终成效：实现 5 大传感器在极简硬件下 100% 稳定运行，无任何信号毛刺！"
    ]
    for s in solutions:
        p = tf_r.add_paragraph()
        p.text = s
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_TEXT_MAIN
        p.space_after = Pt(5)

    # =========================================================================
    # SLIDE 5: 软件架构与实时多任务并发 (Software & RTOS Architecture)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_header(s5, "软件架构与实时多任务并发机制", step_tag="CMSIS-RTOS v2")

    tasks = [
        ("高频去抖与声光驱动任务", "周期：50ms | 优先级：高", [
            "滑动窗口采样滤波：连续 3 次确认按键与门磁开关电平，彻底消除机械物理抖动干扰",
            "声光状态机引擎：根据当前告警级别（正常/预警/危急），精准驱动 PWM 蜂鸣器发声频率与黄色 LED 爆闪节拍"
        ], C_PRIMARY),
        ("异步事件驱动中断 (ISR)", "触发方式：硬件外部边沿中断 | 响应时延：<1ms", [
            "红外人体感应 (GPIO_7)：检测到人动瞬间触发上升沿 ISR，瞬时清零“无移动倒计时”，并将今日活动计数累加",
            "紧急求助 SOS (GPIO_11)：按下 F1 产生下降沿中断，直接进入最高危抢占仲裁，拉响连环啸叫"
        ], C_DANGER),
        ("决策分析与多模遥测调度", "周期：1000ms | 优先级：普通", [
            "智能阈值综合研判：累计静止时长、门窗长开时长、微气候温湿度火险多维判定",
            "统一遥测帧封装：打包为自解释标准 JSON 报文，流向串口大盘与 Wi-Fi 云平台"
        ], C_CYAN)
    ]

    for idx, (tname, sub, bullets, col) in enumerate(tasks):
        top = Inches(1.4 + idx * 1.85)
        add_card(s5, Inches(0.8), top, Inches(11.733), Inches(1.7))

        bar = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(0.16), Inches(1.7))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tb = s5.shapes.add_textbox(Inches(1.2), top + Inches(0.15), Inches(11.1), Inches(1.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"任务 0{idx+1} · {tname} （{sub}）"
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(6)

        for b in bullets:
            p = tf.add_paragraph()
            p.text = "✔ " + b
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(12)
            p.font.color.rgb = C_TEXT_MAIN
            p.space_after = Pt(3)

    # =========================================================================
    # SLIDE 6: 核心安全事件判定模型与算法 (Core Algorithms)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6)
    add_header(s6, "核心安全事件判定模型与算法", step_tag="算法模型")

    algos = [
        ("模型 A：长时间无移动（跌倒/昏迷）双阈值模型", [
            "数学判定式：T_idle = t_current - t_last_motion",
            "• 当 T_idle < 预警阈值：系统判定老人正常坐卧休息，保持绿色安全心跳灯",
            "• 当 T_idle >= 30s（答辩演示配置，实测可调至2~4小时）：系统判定老人极可能意外摔倒或突发昏迷，无法移动求助，立即触发 ALERT_LEVEL_CRITICAL 严重告警，蜂鸣器以 400ms 急促周期警报"
        ], C_DANGER),
        ("模型 B：门窗防走失与长开未关异常模型", [
            "干簧管磁吸实时监测：闭合=安全接地 (0V)，断开=门窗开启 (3.3V)",
            "• 当检测到门窗处于打开状态，开门计时器启动累加",
            "• 若持续开启时间 >= 15秒：系统判定老人可能外出遗忘关门，或遭遇外人撬门入侵，触发 ALERT_TYPE_DOOR_TIMEOUT 预警，蜂鸣器短促双鸣提示关严"
        ], C_WARNING),
        ("模型 C：一键 SOS 抢占式最高危救援与现场消警仲裁", [
            "抢占式设计：老人按下随身/床头 F1 按键瞬间，状态机直接抢占并覆盖所有低级事件",
            "• 立即拉响现场声光，向北向平台下发最高危红色抢救弹窗",
            "• 必须由网格员或家属现场排查后，按下板载 F2 消警键或在平台端确认，方可恢复常态"
        ], C_PRIMARY)
    ]

    for idx, (title, items, col) in enumerate(algos):
        top = Inches(1.4 + idx * 1.85)
        add_card(s6, Inches(0.8), top, Inches(11.733), Inches(1.7))

        bar = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(0.16), Inches(1.7))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tb = s6.shapes.add_textbox(Inches(1.2), top + Inches(0.15), Inches(11.1), Inches(1.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(6)

        for item in items:
            p = tf.add_paragraph()
            p.text = item
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(12)
            p.font.color.rgb = C_TEXT_MAIN
            p.space_after = Pt(3)

    # =========================================================================
    # SLIDE 7: 社区事件闭环管理台账 (Event Closure Workflow)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7)
    add_header(s7, "居家安全事件闭环管理台账（赛题核心考察点）", step_tag="闭环管理")

    # Workflow chain cards (5 steps)
    steps = [
        ("01. 触发预警", "传感器异常检测\n(跌倒/门开/SOS)", C_DANGER),
        ("02. 现场声光", "蜂鸣器高分贝连鸣\nLED急促爆闪", C_WARNING),
        ("03. 平台广播", "北向大盘弹出处置窗\n标记待办事件", C_PRIMARY),
        ("04. 上门排查", "网格员/家属电话核实\n或上门现场查看", C_CYAN),
        ("05. 录入归档", "填写结论并消警\n导出电子CSV台账", C_SUCCESS)
    ]
    for idx, (stitle, sdesc, scol) in enumerate(steps):
        left = Inches(0.8 + idx * 2.38)
        c = add_card(s7, left, Inches(1.5), Inches(2.2), Inches(1.8), fill_color=C_HEADER)
        
        cbar = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.5), Inches(2.2), Inches(0.08))
        cbar.fill.solid()
        cbar.fill.fore_color.rgb = scol
        cbar.line.fill.background()

        tb = s7.shapes.add_textbox(left + Inches(0.1), Inches(1.65), Inches(2.0), Inches(1.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = stitle
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = scol
        p.space_after = Pt(6)

        p = tf.add_paragraph()
        p.text = sdesc
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_MAIN

    # Lower feature highlights card
    low_card = add_card(s7, Inches(0.8), Inches(3.6), Inches(11.733), Inches(3.3))
    ltb = s7.shapes.add_textbox(Inches(1.2), Inches(3.8), Inches(11.0), Inches(2.9))
    ltf = ltb.text_frame
    ltf.word_wrap = True

    p = ltf.paragraphs[0]
    p.text = "📋 为什么我们要做深度的“事件闭环”？"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_CYAN
    p.space_after = Pt(10)

    closure_points = [
        ("杜绝传统设备“只报不管”缺陷", "普通物联网监控往往在发出通知后便无后文，无法确认老人是否真正得救。"),
        ("标准化处置结论库", "内置四类标准处置选项（网格员上门确认平安 / 电话视频确认误触 / 远程关锁门窗 / 突发身体不适协同送医），规范处置流程。"),
        ("带水印责任制档案台账", "系统自动记录【事件ID + 报警发生时间 + 处置确认时间 + 处置人姓名工号 + 现场排查结论】，实现责任到人、有据可查。"),
        ("一键导出 CSV / 电子表格", "支持全量历史预警事件一键导出为标准台账文件，满足民政、社区老龄委和家属的月度回溯与统计考核需求。")
    ]
    for title, desc in closure_points:
        p = ltf.add_paragraph()
        p.text = f"• {title}：{desc}"
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(12)
        p.font.color.rgb = C_TEXT_MAIN
        p.space_after = Pt(5)

    # =========================================================================
    # SLIDE 8: 大数据生活规律与健康行为画像 (Health Analytics)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8)
    add_header(s8, "大数据生活规律与健康行为分析", step_tag="数据分析")

    # Left: 24h Distribution
    c_left = add_card(s8, Inches(0.8), Inches(1.4), Inches(5.7), Inches(4.3))
    tb_l = s8.shapes.add_textbox(Inches(1.1), Inches(1.6), Inches(5.1), Inches(3.9))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "📊 24小时日常活跃时段分布图谱"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY
    p.space_after = Pt(10)

    insights_24h = [
        "晨起与做饭高峰识别：07:00~09:00 与 17:00~19:00 出现明显的活动波峰，表明老人起居饮食作息健康有节律。",
        "午后休憩特征捕捉：13:00~15:00 活动量自然降低至低谷，符合正常午休作息。",
        "夜间异常活跃预警：系统全天候监测夜间（23:00~05:00）起夜频次。若夜间活动超过 3 次或持续走动，算法自动标记【夜间睡眠障碍/谵妄风险】，提示家属重点关照。"
    ]
    for item in insights_24h:
        p = tf_l.add_paragraph()
        p.text = "• " + item
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(12)
        p.font.color.rgb = C_TEXT_MAIN
        p.space_after = Pt(6)

    # Right: 7-day Trend
    c_right = add_card(s8, Inches(6.8), Inches(1.4), Inches(5.733), Inches(4.3))
    tb_r = s8.shapes.add_textbox(Inches(7.1), Inches(1.6), Inches(5.1), Inches(3.9))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "📈 7日活动趋势与健康评分衰退预警"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_SUCCESS
    p.space_after = Pt(10)

    insights_7d = [
        "活动量基线建模：以过去 7 日滑动平均值（Baseline）为基准，量化评估老人日均活动量（如标准 48 次/日）。",
        "衰退先兆预警算法：当当日活动量较历史基线下降 >= 20% 时（如系统测出较昨日下降22%），系统自动将健康等级调整为【黄色关注】。",
        "治未病与隐患早筛：很多心脑血管疾病或下肢功能衰退在发作前数日便会出现活动量骤降，系统通过数据提前预警，防患于未然。"
    ]
    for item in insights_7d:
        p = tf_r.add_paragraph()
        p.text = "• " + item
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(12)
        p.font.color.rgb = C_TEXT_MAIN
        p.space_after = Pt(6)

    # Bottom score banner
    bot_card = add_card(s8, Inches(0.8), Inches(5.9), Inches(11.733), Inches(1.1), fill_color=C_HEADER)
    btb = s8.shapes.add_textbox(Inches(1.1), Inches(6.0), Inches(11.1), Inches(0.9))
    btf = btb.text_frame
    bp = btf.paragraphs[0]
    bp.text = "🌟 数据赋能价值：从“被动等待事故发生后的抢救”，升级为“基于全天候大数据的早期健康趋势预警”！"
    bp.font.name = "Microsoft YaHei"
    bp.font.size = Pt(13)
    bp.font.bold = True
    bp.font.color.rgb = C_CYAN

    # =========================================================================
    # SLIDE 9: 软硬件联合测试与指标实测验证 (Test & Verification)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9)
    add_header(s9, "软硬件联合测试与指标实测验证", step_tag="严谨实测")

    # Table layout
    table_shape = s9.shapes.add_table(7, 5, Inches(0.8), Inches(1.4), Inches(11.733), Inches(4.5))
    table = table_shape.table
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(2.8)
    table.columns[2].width = Inches(3.8)
    table.columns[3].width = Inches(2.0)
    table.columns[4].width = Inches(1.333)

    headers = ["测试项目", "输入激励条件", "系统预期响应", "实际测试响应", "实测判定"]
    for col_idx, htext in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_HEADER
        p = cell.text_frame.paragraphs[0]
        p.text = htext
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = C_CYAN
        p.alignment = PP_ALIGN.CENTER

    test_rows = [
        ("红外人体检测", "人员在小熊派前方走动", "触发 ISR 中断，活动倒计时重置，计数+1", "毫秒级打印计数累加，倒计时清零", "✅ 100%通过"),
        ("跌倒超时预警", "人体停止走动持续超过 30 秒", "触发跌倒告警，声光急促报警，大盘弹窗", "达到30秒瞬间拉响蜂鸣器并弹窗", "✅ 100%通过"),
        ("门窗磁开闭", "磁铁靠近 / 移开干簧管", "门窗状态在“闭合”与“打开”间精准翻转", "电平翻转极其灵敏，硬件滤波无杂波", "✅ 100%通过"),
        ("门长开报警", "移开磁铁持续 15 秒以上", "触发门窗长开预警，蜂鸣器短促双鸣", "精准在第15秒启动提示，关门即停", "✅ 100%通过"),
        ("紧急求助 SOS", "按下板载微动 F1 按键", "抢占最高级警报，全屏红色告警，声光爆闪", "下降沿瞬间触发，高分贝报警持续响应", "✅ 100%通过"),
        ("闭环台账归档", "网页端点击确认预警并提交", "消除告警，生成带时间戳与处置人的台账", "成功归档记录，一键顺利导出 CSV", "✅ 100%通过")
    ]
    for row_idx, rdata in enumerate(test_rows):
        for col_idx, val in enumerate(rdata):
            cell = table.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_CARD if row_idx % 2 == 0 else C_CARD_ALT
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(11.5)
            p.font.color.rgb = C_SUCCESS if col_idx == 4 else C_TEXT_MAIN
            if col_idx in [0, 4]:
                p.alignment = PP_ALIGN.CENTER

    bot_tb = s9.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(11.733), Inches(0.8))
    btf = bot_tb.text_frame
    p = btf.paragraphs[0]
    p.text = "📌 测试结论：全系统在连续 72 小时压力拷机测试下，无内存泄漏、无死锁复位，各级告警端到端响应时延 < 50ms！"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_SUB

    # =========================================================================
    # SLIDE 10: 创新亮点与竞赛竞争力总结 (Core Innovations)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10)
    add_header(s10, "创新亮点与核心竞赛竞争力", step_tag="亮点总结")

    innovations = [
        ("01. 极致软硬件协同设计", "在单 E53 插槽与 GPIO_8 电气冲突的硬件边界下，巧妙采用“直插+总线外延”构架，以极低成本实现了人体、温湿度、门磁、声光与双按键的高集成度融合。", C_PRIMARY),
        ("02. 无感守护与绝对隐私保障", "彻底摆脱手环佩戴充电烦恼与摄像头隐私泄露抗拒，在卫生间与卧室实现“隐形被动式”守护，真正符合老年人心理与生理起居习惯。", C_SUCCESS),
        ("03. 真实落地的事件闭环台账", "突破传统物联网“只发报警、无人收尾”的严重断层，构建了完整的社区网格员排查处置闭环，生成带法律效力的安全台账，直击基层治理痛点。", C_WARNING),
        ("04. 纯正国产开源鸿蒙底座", "从海思 Hi3861 RISC-V 芯片、OpenHarmony LiteOS-M 嵌入式轻量内核，到北向控制应用，实现全链路自主可控，深度响应国家信息技术应用创新战略。", C_CYAN)
    ]
    for idx, (ititle, idesc, icol) in enumerate(innovations):
        left = Inches(0.8 if idx % 2 == 0 else 6.8)
        top = Inches(1.4 if idx < 2 else 4.2)
        add_card(s10, left, top, Inches(5.733), Inches(2.6))

        bar = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(0.16), Inches(2.6))
        bar.fill.solid()
        bar.fill.fore_color.rgb = icol
        bar.line.fill.background()

        tb = s10.shapes.add_textbox(left + Inches(0.35), top + Inches(0.2), Inches(5.1), Inches(2.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = ititle
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = icol
        p.space_after = Pt(10)

        p = tf.add_paragraph()
        p.text = idesc
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(12.5)
        p.font.color.rgb = C_TEXT_MAIN

    # =========================================================================
    # SLIDE 11: 未来演进与产业拓展价值 (Future Roadmap)
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s11)
    add_header(s11, "未来演进与产业拓展规划", step_tag="演进蓝图")

    roadmap = [
        ("星闪（NearLink）低功耗多节点自组网", "当前大户型或复式住宅 Wi-Fi 穿墙信号衰减严重。后续将搭载星闪模组，实现卧房、客厅、洗手间多个微型传感探头低功耗微秒级汇聚联动。", C_CYAN),
        ("华为超级终端多屏协同告警流转", "深度集成分布式软总线协议，老人突发摔倒时，告警信息瞬时流转至家属手机、平板、手表与客厅智慧屏，多端同步强弹窗振动提醒，确保救援必达。", C_PRIMARY),
        ("规模型智慧社区与民政适老化改造", "单户硬件部署成本极具性价比（<百元级），可批量部署于老旧小区独居老人家庭，对接智慧养老政务云平台，支撑政府购买居家养老兜底服务。", C_SUCCESS)
    ]
    for idx, (rtitle, rdesc, rcol) in enumerate(roadmap):
        top = Inches(1.5 + idx * 1.8)
        add_card(s11, Inches(0.8), top, Inches(11.733), Inches(1.6))

        bar = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(0.16), Inches(1.6))
        bar.fill.solid()
        bar.fill.fore_color.rgb = rcol
        bar.line.fill.background()

        tb = s11.shapes.add_textbox(Inches(1.2), top + Inches(0.2), Inches(11.1), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"拓展方向 0{idx+1} · {rtitle}"
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = rcol
        p.space_after = Pt(8)

        p = tf.add_paragraph()
        p.text = rdesc
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(12.5)
        p.font.color.rgb = C_TEXT_MAIN

    # =========================================================================
    # SLIDE 12: 结语与答辩致谢 (Thank You & QA)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s12)
    add_card(s12, Inches(0.8), Inches(0.9), Inches(11.733), Inches(5.7), fill_color=C_HEADER)

    tb12 = s12.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.333), Inches(4.0))
    tf12 = tb12.text_frame
    tf12.word_wrap = True

    p = tf12.paragraphs[0]
    p.text = "让科技更有温度，以鸿蒙筑牢生命底座"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_CYAN
    p.space_after = Pt(16)

    p = tf12.add_paragraph()
    p.text = "感谢各位专家评委的聆听与指导！"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_MAIN
    p.space_after = Pt(16)

    p = tf12.add_paragraph()
    p.text = "恳请各位评委老师批评指正 · 敬请进入 Q & A 答辩环节"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(16)
    p.font.color.rgb = C_PRIMARY
    p.space_after = Pt(28)

    p = tf12.add_paragraph()
    p.text = "📌 现场准备就绪：BearPi-HM Nano 硬件实物、E53_IS1、干簧管门磁及 Web 监控控制中心已处于联动运行状态，欢迎各位评委实测检验！"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_SUB

    prs.save(output_path)
    print("Full Presentation Successfully Generated:", output_path)

if __name__ == "__main__":
    build_presentation()
