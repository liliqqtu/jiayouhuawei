# -*- coding: utf-8 -*-
"""给路演 PPT 补 3 页（分级播报 / 睡眠抑制 / 软硬件合流路径），并重排序号。

设计令牌全部沿用原文件：
  标题 28pt bold #0D1326   Noto Sans SC   (1.11, 0.83) 11.11 x 0.72
  强调条 0.56 x 0.04 #2563EB (1.11, 1.71)
  副标题 14pt #64748B (1.11, 1.95) 11.11 x 0.36
  白卡片 FFFFFF / 边线 DCE3EC，深色卡片 0D1326
  等级色 L3 B91C1C  L2 EA580C  L1 60A5FA
  页脚 9pt #94A3B8 (1.11, 7.02)；页码 (11.9, 7.02)

用法：python scripts/ppt_add_pages.py [--in 文件] [--check]
"""
from __future__ import print_function

import argparse
import os
import shutil
import sys
import tempfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

FONT = "Noto Sans SC"

INK = RGBColor(0x0D, 0x13, 0x26)      # 深墨（标题/卡片标题）
BLUE = RGBColor(0x25, 0x63, 0xEB)     # 品牌蓝（强调条/编号块）
SUB = RGBColor(0x64, 0x74, 0x8B)      # 副标题灰
BODY = RGBColor(0x33, 0x41, 0x55)     # 正文
DIM = RGBColor(0x94, 0xA3, 0xB8)      # 弱化文字
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD_LINE = RGBColor(0xDC, 0xE3, 0xEC)
L3 = RGBColor(0xB9, 0x1C, 0x1C)
L2 = RGBColor(0xEA, 0x58, 0x0C)
L1 = RGBColor(0x60, 0xA5, 0xFA)
DARK_LIGHT = RGBColor(0x94, 0xA3, 0xB8)


# ---------------------------------------------------------------- 基础绘制
def text(slide, x, y, w, h, content, size=11.5, color=BODY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """放一个纯文本框。content 可以是字符串，也可以是 [(文字, {覆盖参数}), ...] 列表。"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor

    runs = content if isinstance(content, list) else [(content, {})]
    first = True
    for value, override in runs:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = override.get("align", align)
        para.line_spacing = override.get("spacing", spacing)
        para.space_before = Pt(0)
        para.space_after = Pt(0)
        run = para.add_run()
        run.text = value
        font = run.font
        font.name = FONT
        font.size = Pt(override.get("size", size))
        font.bold = override.get("bold", bold)
        font.color.rgb = override.get("color", color)
    return box


def block(slide, x, y, w, h, fill=WHITE, line=CARD_LINE, radius=True):
    """圆角/直角色块。"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    if radius:
        shape.adjustments[0] = 0.06
    return shape


def header(slide, title, subtitle, section):
    """统一的页头 + 页脚。"""
    text(slide, 1.11, 0.83, 11.11, 0.72, title, size=28, color=INK, bold=True)
    block(slide, 1.11, 1.71, 0.56, 0.04, fill=BLUE, line=None, radius=False)
    text(slide, 1.11, 1.95, 11.11, 0.36, subtitle, size=14, color=SUB)
    text(slide, 1.11, 7.02, 6.0, 0.28, u"守夜灯 · 赛题四\u3000|\u3000" + section,
         size=9, color=DIM)
    # 页码占位，reorder() 会按最终位置改写
    text(slide, 11.9, 7.02, 0.32, 0.28, u"0", size=9, color=DIM)


def notes(slide, content):
    slide.notes_slide.notes_text_frame.text = content


# ---------------------------------------------------------------- 三页内容
def build_sleep(prs):
    """A：睡眠期抑制 —— 放在「三档 + 融合规则」之后。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, u"夜里不该被吵醒",
           u"睡眠期判断的不是“动得少”，而是“该动没动”——夜间误报是这类系统最容易翻车的地方。",
           u"判定")

    rows = [
        (u"入睡", L1, u"23:00–01:00　走动减少、门磁关闭 = 正常入睡", u"不报警，只记录"),
        (u"深睡", L2, u"01:00–07:00　连续 6 小时零活动才确认异常", u"L2 待确认"),
        (u"离床", L3, u"夜间卫生间 2 分钟无活动 = 最高风险组合", u"L3 立即告警"),
    ]
    for index, (label, color, body, level) in enumerate(rows):
        top = 2.52 + index * 1.28
        block(slide, 1.11, top, 6.85, 1.1)
        block(slide, 1.11, top + 0.12, 0.06, 0.86, fill=color, line=None, radius=False)
        text(slide, 1.39, top + 0.14, 1.3, 0.32, label, size=14, color=color, bold=True)
        text(slide, 2.61, top + 0.12, 3.05, 0.6, body, size=11.5, color=BODY)
        text(slide, 5.79, top + 0.12, 2.05, 0.85, level, size=10.5, color=SUB)

    block(slide, 8.35, 2.35, 3.87, 4.05, fill=INK, line=None)
    text(slide, 8.62, 2.55, 3.35, 0.32, u"降噪三招", size=12, color=L1, bold=True)
    items = [
        (u"① 睡眠窗放宽", u"23:00–07:00，“无活动”阈值从 3 分钟放宽到 360 分钟"),
        (u"② 卫生间单列", u"唯一不受睡眠抑制的区域，反而收紧到 2 分钟"),
        (u"③ 白天型静音", u"夜里只允许 L3 出声，L1 / L2 只记不发"),
    ]
    for index, (head, body) in enumerate(items):
        top = 3.0 + index * 1.12
        text(slide, 8.62, top, 3.35, 0.3, head, size=12.5, color=WHITE, bold=True)
        text(slide, 8.62, top + 0.3, 3.35, 0.72, body, size=10.5, color=DARK_LIGHT)

    notes(slide, u"45 秒：夜里老人不动是正常的。如果沿用白天的 3 分钟无活动阈值，"
                 u"一晚上要报几十次。所以睡眠窗内把“无活动”放宽到 6 小时，"
                 u"只有卫生间例外——起夜是最危险的动作，反而收紧到 2 分钟。"
                 u"这条是我们把误报压下来的关键。")
    return slide


def build_siren(prs):
    """B：三级播报 —— 放在「呼救四段接力」之后。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, u"声音也要分级",
           u"同样的告警，半夜的 L3 和下午的 L1 不该是同一种声音——网关本机分级播报，断网也照响。",
           u"响应")

    rows = [
        (u"L3 紧急", L3, u"扫频警笛循环 600→1600Hz　约每 10 秒一轮，直到有人确认",
         u"语音：紧急告警……请立即处理"),
        (u"L2 关注", L2, u"三声中速提示（1200 / 1200 / 950Hz）　每 10 秒一次",
         u"语音：注意……请尽快确认"),
        (u"L1 提醒", L1, u"两声轻提示（800Hz）　每 15 秒一次",
         u"语音：提醒……"),
    ]
    for index, (label, color, body, voice) in enumerate(rows):
        top = 2.52 + index * 1.28
        block(slide, 1.11, top, 6.85, 1.1)
        block(slide, 1.11, top + 0.12, 0.06, 0.86, fill=color, line=None, radius=False)
        text(slide, 1.39, top + 0.14, 1.3, 0.32, label, size=14, color=color, bold=True)
        text(slide, 2.61, top + 0.12, 5.1, 0.5, body, size=11, color=BODY)
        text(slide, 2.61, top + 0.62, 5.1, 0.3, voice, size=10, color=SUB)

    block(slide, 8.35, 2.35, 3.87, 4.05, fill=INK, line=None)
    text(slide, 8.62, 2.55, 3.35, 0.32, u"为什么单做一页", size=12, color=L1, bold=True)
    items = [
        (u"① 分级决定打扰", u"L1 只提示，L3 能把人叫醒"),
        (u"② 按级轮转", u"多条待确认时 L3→L2→L1 排序，最严重的先报"),
        (u"③ 确认即停", u"全部处置完声音立刻停；未确认则一直响"),
    ]
    for index, (head, body) in enumerate(items):
        top = 3.0 + index * 1.12
        text(slide, 8.62, top, 3.35, 0.3, head, size=12.5, color=WHITE, bold=True)
        text(slide, 8.62, top + 0.3, 3.35, 0.72, body, size=10.5, color=DARK_LIGHT)

    notes(slide, u"40 秒：这一页讲“吵人的分寸”。三级不是设计好看，是夜里真不能一视同仁。"
                 u"L3 是循环扫频，能把人叫醒；L1 只是两声轻响。"
                 u"多条同时来时按等级排序轮转，最严重的先报，确认完就停。")
    return slide


def build_merge(prs):
    """C：软硬件合流路径 —— 放在「工程与交付」之后。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, u"两条线怎么合成一条？",
           u"固件与 APP 并行开发，靠“先定契约、再做等价替换”保证最后一定能合上。",
           u"工程")

    steps = [
        (u"01", u"契约先行", u"字段 16 个、事件 10 类、端口 8780 / 9998 一次定死"),
        (u"02", u"模拟器等价", u"device_sim 与真板同协议、同字段，先把全链路跑通"),
        (u"03", u"两侧并行", u"南向做板端判定与上报；北向 APP 只认契约，不认实现"),
        (u"04", u"合流入模", u"模拟数据与真机数据进同一套判定、同一个界面"),
    ]
    for index, (num, head, body) in enumerate(steps):
        top = 2.45 + index * 1.02
        block(slide, 1.11, top, 6.85, 0.86)
        small = block(slide, 1.39, top + 0.25, 0.52, 0.36, fill=BLUE, line=None)
        small.text_frame.word_wrap = False
        para = small.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = num
        run.font.name = FONT
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = WHITE
        text(slide, 2.09, top + 0.21, 1.5, 0.32, head, size=14, color=INK, bold=True)
        text(slide, 3.55, top + 0.22, 4.2, 0.5, body, size=10.5, color=BODY)

    block(slide, 8.35, 2.35, 3.87, 4.05, fill=INK, line=None)
    text(slide, 8.62, 2.55, 3.35, 0.32, u"合流检查清单", size=12, color=L1, bold=True)
    items = [
        (u"契约一致", u"check_firmware.py 静态核对字段、宏、端口"),
        (u"阈值同值", u"config.h 与 config.py 逐项对齐，改一处必须改两处"),
        (u"上板三步", u"烧录 → hdc 连上网关 → 断网测本地判定"),
        (u"风险兜底", u"板子出问题时模拟器顶着演示，不空场"),
    ]
    for index, (head, body) in enumerate(items):
        top = 3.0 + index * 0.86
        text(slide, 8.62, top, 3.35, 0.28, head, size=12, color=WHITE, bold=True)
        text(slide, 8.62, top + 0.28, 3.35, 0.52, body, size=10, color=DARK_LIGHT)

    notes(slide, u"35 秒：这部分最容易被问“你们两条线怎么保证能合上”。"
                 u"答案是契约先行：字段、事件类型、端口先写死在文档里，两边各写各的；"
                 u"模拟器跟真板走同一套协议，所以现在软件侧每一条验证，换到真板上都成立。"
                 u"真板出问题也不影响现场，模拟器可以顶。")
    return slide


# ---------------------------------------------------------------- 重排与页码
def reorder(prs, extra):
    """把 3 个新页插到指定位置，并重编全部页码。

    新页是最后追加的，所以 sldIdLst 末尾三个元素依次是 sleep / siren / merge。
    最终顺序：1–8 | 睡眠 | 9 | 播报 | 10–17 | 合流 | 18
    """
    sld_ids = prs.slides._sldIdLst
    elements = list(sld_ids)
    if len(elements) != 21:
        raise RuntimeError("期望 21 页，实际 %d" % len(elements))

    old = elements[:18]
    sleep_el, siren_el, merge_el = elements[18], elements[19], elements[20]
    new_order = (old[0:8] + [sleep_el, old[8]] + [siren_el] +
                 old[9:17] + [merge_el, old[17]])

    # lxml 的 append 会自动把元素从原位置移到末尾，逐个子即可完成重排
    for element in new_order:
        sld_ids.append(element)

    # 页码：找右下角那个纯数字小文本框
    for index, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            try:
                left = shape.left / 914400.0
                top = shape.top / 914400.0
            except TypeError:
                continue
            if left > 11.6 and top > 6.9:
                frame = shape.text_frame.text.strip()
                if frame.isdigit():
                    shape.text_frame.paragraphs[0].runs[0].text = str(index + 1)
    return new_order


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    # PPT 放在工作区根目录（software 的上一级）
    root = os.path.dirname(os.path.dirname(here))
    default = os.path.join(root, u"赛题四_路演PPT_10分钟.pptx")
    parser = argparse.ArgumentParser()
    parser.add_argument("--in", dest="src", default=default)
    parser.add_argument("--check", action="store_true", help=u"只统计不写文件")
    args = parser.parse_args()

    prs = Presentation(args.src)
    print(u"读入：%s（%d 页）" % (os.path.basename(args.src), len(prs.slides)))

    if args.check:
        # 重复执行保护：已经补过就不要再补
        titles = set()
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    titles.add(shape.text_frame.text.strip())
        for key in (u"夜里不该被吵醒", u"声音也要分级", u"两条线怎么合成一条？"):
            print(u"  含「%s」：%s" % (key, key in titles))
        return 0

    if len(prs.slides) != 18:
        print(u"[中止] 页数不是 18（当前 %d），可能已经补过了。" % len(prs.slides))
        return 2

    backup = os.path.join(tempfile.gettempdir(),
                          u"路演PPT_18页_备份_%d.pptx" % os.getpid())
    shutil.copyfile(args.src, backup)
    print(u"已备份到：%s" % backup)

    extra = {
        "sleep": build_sleep(prs),
        "siren": build_siren(prs),
        "merge": build_merge(prs),
    }
    order = reorder(prs, extra)
    prs.save(args.src)
    print(u"完成：%d 页 → %s" % (len(order), args.src))
    for index, slide in enumerate(prs.slides):
        first = ""
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                first = shape.text_frame.text.strip().splitlines()[0]
                break
        print(u"  %02d  %s" % (index + 1, first[:40]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
